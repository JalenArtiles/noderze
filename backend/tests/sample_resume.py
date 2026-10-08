"""Builds a fictional resume (.docx) for tests, so the suite never needs a real person's resume.

It deliberately contains the three problems the resume checker should catch:
an education end date before the expected graduation, a sales-engineering overclaim, and a course-name typo.
"""
from __future__ import annotations

from pathlib import Path

import docx


def _heading(d, text: str) -> None:
    d.add_paragraph().add_run(text).bold = True


def _role(d, head: str, dates: str) -> None:
    p = d.add_paragraph()
    p.add_run(head).bold = True
    p.add_run("\t" + dates)


def build(path: str | Path) -> Path:
    d = docx.Document()
    d.add_paragraph().add_run("SAMPLE CANDIDATE").bold = True
    d.add_paragraph("Tempe, AZ | sample.candidate@example.com | 555-010-0199 | linkedin.com/in/sample-candidate")

    _heading(d, "PROFESSIONAL SUMMARY")
    d.add_paragraph("Cybersecurity student experienced in sales engineering and technical discovery, "
                    "with B2B prospecting and cold calling experience.")

    _heading(d, "EDUCATION")
    _role(d, "B.S. Applied Computing (Cybersecurity) - Arizona State University", "2023 - 2026")

    _heading(d, "RELEVANT COURSEWORK")
    d.add_paragraph("Netwrok Security, Operating Systems, Database Systems, Computer Networks")

    _heading(d, "EXPERIENCE")
    _role(d, "Sales Development Intern - Example Networks", "May 2025 - Aug 2025")
    d.add_paragraph("Phoenix, AZ")
    for line in ("Made cold calls to IT and security leaders and booked discovery meetings for Account Executives.",
                 "Prospected mid-market accounts and kept CRM records current in HubSpot.",
                 "Translated technical capabilities of network security products into business value for prospects."):
        d.add_paragraph(line, style="List Bullet")

    _heading(d, "PROJECTS")
    _role(d, "Home Lab - Personal Project", "Jan 2025 - Present")
    d.add_paragraph("Built a small home lab with a VM firewall and wrote a one-page setup guide.", style="List Bullet")

    _heading(d, "SKILLS")
    d.add_paragraph("Python, Java, SQL, Bash, Wireshark, Computer networks, Network security, CRM, Prospecting, Discovery")

    path = Path(path)
    d.save(str(path))
    return path
